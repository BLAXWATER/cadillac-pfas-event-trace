"""Read-only content previews. Text snapshots are explicitly not original pagination."""
import hashlib, json, re, subprocess, sys, textwrap, zipfile, struct
from pathlib import Path
from urllib.parse import urlparse, unquote
from xml.etree import ElementTree
from html.parser import HTMLParser
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tmp/preview-deps'))
import olefile
from openpyxl import load_workbook

def key(url): return '/'+unquote(urlparse(url).path).split('/public/')[-1].lstrip('/')
def original(row):
    p=key(row['url'])
    for base in [ROOT/'public']+[d/'public' for d in ROOT.parent.iterdir() if (d/'public').is_dir()]:
        f=base/p.lstrip('/')
        if f.is_file() and hashlib.sha256(f.read_bytes()).hexdigest()==row['sha256']: return f
    match=re.search(r'/blob/([a-f0-9]{40})/',row['url'])
    if match:
        r=subprocess.run(['git','show',match[1]+':public'+p],cwd=ROOT,capture_output=True)
        if r.returncode==0 and hashlib.sha256(r.stdout).hexdigest()==row['sha256']:
            f=ROOT/'tmp/preview-originals'/p.lstrip('/');f.parent.mkdir(parents=True,exist_ok=True);f.write_bytes(r.stdout);return f
        from urllib.request import urlopen
        url=row['url'].replace('https://github.com/cazey43/','https://raw.githubusercontent.com/BLAXWATER/').replace('https://github.com/BLAXWATER/','https://raw.githubusercontent.com/BLAXWATER/').replace('/blob/','/')
        data=urlopen(url,timeout=60).read()
        if hashlib.sha256(data).hexdigest()==row['sha256']:
            f=ROOT/'tmp/preview-originals'/p.lstrip('/');f.parent.mkdir(parents=True,exist_ok=True);f.write_bytes(data);return f
    raise ValueError('Original unavailable or hash mismatch: '+row['name'])

class TextHTML(HTMLParser):
    def __init__(self): super().__init__();self.parts=[];self.hidden=0
    def handle_starttag(self,tag,attrs):
        if tag in ['script','style']:self.hidden+=1
    def handle_endtag(self,tag):
        if tag in ['script','style']:self.hidden=max(0,self.hidden-1)
    def handle_data(self,data):
        if not self.hidden and data.strip():self.parts.append(data.strip())

def content(f):
    ext=f.suffix.lower()
    if ext=='.xlsx':
        wb=load_workbook(f,read_only=True,data_only=False);s=wb.worksheets[0]
        lines=['Worksheet: '+s.title+' (cell values/formulas, not print layout)']
        for row in s.iter_rows(min_row=1,max_row=min(s.max_row or 35,35),max_col=min(s.max_column or 20,20)):
            lines.append(' | '.join(c.coordinate+'='+str(c.value) for c in row if c.value is not None))
        wb.close();return '\n'.join(lines)
    if ext=='.docx':
        with zipfile.ZipFile(f) as z:r=ElementTree.fromstring(z.read('word/document.xml'))
        return '\n'.join(''.join(p.itertext()) for p in r.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'))
    if ext=='.msg':
        with olefile.OleFileIO(f) as o:
            def prop(n):
                for suffix,encoding in [('001F','utf-16le'),('001E','cp1252')]:
                    p='__substg1.0_'+n+suffix
                    if o.exists(p):return o.openstream(p).read().decode(encoding,errors='replace').strip('\x00')
                return ''
            return 'Subject: '+prop('0037')+'\nFrom: '+prop('0C1A')+'\nTo: '+prop('0E04')+'\n\n'+prop('1000')
    if ext=='.doc':
        with olefile.OleFileIO(f) as o:
            d=o.openstream('WordDocument').read(); t=o.openstream('1Table' if struct.unpack_from('<H',d,10)[0]&0x200 else '0Table').read()
            fc,size=struct.unpack_from('<II',d,0x1a2);clx=t[fc:fc+size];i=0
            while clx[i]==1:i+=3+struct.unpack_from('<H',clx,i+1)[0]
            if clx[i]!=2:raise ValueError('Unsupported Word piece table')
            size=struct.unpack_from('<I',clx,i+1)[0];plc=clx[i+5:i+5+size];n=(size-4)//12;parts=[]
            for j in range(n):
                start,end=struct.unpack_from('<II',plc,j*4);offset=struct.unpack_from('<I',plc,(n+1)*4+j*8+2)[0];compressed=bool(offset&0x40000000);offset &=0x3fffffff
                if compressed:offset//=2
                parts.append(d[offset:offset+(end-start)*(1 if compressed else 2)].decode('cp1252' if compressed else 'utf-16le',errors='replace'))
            return ''.join(parts).replace('\r','\n').replace('\x07',' | ')
    if ext=='.zip':
        with zipfile.ZipFile(f) as z:return 'Archive contents (filenames and uncompressed byte sizes)\n'+'\n'.join(f'{x.filename} ({x.file_size} bytes)' for x in z.infolist())
    text=f.read_text(encoding='utf-8-sig',errors='replace')
    if ext=='.html':
        p=TextHTML();p.feed(text);return '\n'.join(p.parts)
    return text

def main():
    manifest={};errors=[]
    font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',22)
    titlefont=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)
    rows=[r for f in (ROOT/'app').glob('*-documents.json') for r in json.loads(f.read_text(encoding='utf-8'))]
    for row in rows:
        if key(row['url']).lower().endswith('.pdf'):continue
        try:
            f=original(row);out=ROOT/'public/first-page-previews/nonpdf'/f'{row["sha256"]}.webp';out.parent.mkdir(parents=True,exist_ok=True)
            if f.suffix.lower() in ['.png','.jpg','.jpeg','.webp']:
                im=Image.open(f).convert('RGB');im.thumbnail((1400,2000));label='Original source image'
            else:
                text=content(f)
                if not text.strip():raise ValueError('No extractable content')
                lines=[]
                for line in text.splitlines():lines.extend(textwrap.wrap(''.join(c for c in line if c.isprintable()),width=96) or [''])
                im=Image.new('RGB',(1400,1800),'white');draw=ImageDraw.Draw(im)
                label='Content excerpt from original file (not original page layout)'
                draw.text((40,25),label,font=titlefont,fill='black')
                for i,line in enumerate(textwrap.wrap(row['name'],95)[:3]):draw.text((40,65+i*29),line,font=font,fill='black')
                for i,line in enumerate(lines[:51]):draw.text((40,175+i*29),line,font=font,fill='black')
                draw.text((40,1715),'Preview excerpt only. Download preserves the complete original file.',font=font,fill='black')
            im.save(out,'WEBP',quality=85)
            manifest[key(row['url'])]={'preview':'/'+out.relative_to(ROOT/'public').as_posix(),'label':label,'sourceSha256':row['sha256']}
        except Exception as e:errors.append({'name':row['name'],'error':str(e)})
    (ROOT/'app/nonpdf-preview-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print('Non-PDF previews:',len(manifest),'Errors:',json.dumps(errors))
    if errors:sys.exit(1)
if __name__=='__main__':main()
